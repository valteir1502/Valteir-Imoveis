import docx

doc_path = r"D:\Negocios Imobiliários\Condomínio Buona Vita\Reginaldo\Contrato_Compromisso_Compra_e_Venda_Ipe_da_Mata.docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

with open("contrato_original_texto.txt", "w", encoding="utf-8") as f:
    for i, p in enumerate(doc.paragraphs):
        f.write(f"[{i}] {p.text}\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\n--- TABLE {t_idx} ---\n")
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.strip() for c in row.cells]
            f.write(f"Row {r_idx}: {' | '.join(cells)}\n")

print("Saved to contrato_original_texto.txt")
