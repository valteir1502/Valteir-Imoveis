import os
import shutil
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_contract():
    doc = Document()

    # Configuração de Margens
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Estilo Normal
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.15, bold=False):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            run = p.add_run(text)
            run.bold = bold
        return p

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return p

    def add_clause_header(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        return p

    # Cabeçalho do Contrato
    add_title("INSTRUMENTO PARTICULAR DE COMPROMISSO DE COMPRA E VENDA DE IMÓVEL COM DAÇÃO EM PAGAMENTO / PERMUTA")
    add_subtitle("VALOR TOTAL DO NEGÓCIO: R$ 2.390.000,00 (DOIS MILHÕES, TREZENTOS E NOVENTA MIL REAIS)")

    # Preâmbulo
    add_p("Neste Instrumento, irrevogável e irretratável, passível de Adjudicação Compulsória nos termos do art. 22 (com a redação dada pela Lei nº 649/1949 e Lei nº 14.382/2022) do Decreto-Lei nº 58, de 10 de Dezembro de 1937 e Código Civil Brasileiro, independente do prévio registro imobiliário, celebrado na presença das 02 (duas) testemunhas ao final indicadas e assinadas, tendo como foro competente a Comarca de São José do Rio Preto - SP, com renúncia expressa a qualquer outro por mais privilegiado que seja, comparecem as partes a seguir qualificadas:")

    # VENDEDORES
    add_clause_header("A – PROMITENTES VENDEDORES:")
    
    p_v1 = add_p()
    r = p_v1.add_run("1. ADEMIR VENTURINI")
    r.bold = True
    p_v1.add_run(", brasileiro, marceneiro, portador da Cédula de Identidade RG nº 24.694.860 SSP/SP e inscrito no CPF/MF sob o nº 133.470.468-62, casado sob o regime da Comunhão Parcial de Bens na vigência da Lei nº 6.515/77 com ")
    r = p_v1.add_run("MARILAINE OLIVEIRA SILVA VENTURINI")
    r.bold = True
    p_v1.add_run(", brasileira, do lar, portadora da Cédula de Identidade RG nº 25.067.901 SSP/SP e inscrita no CPF/MF sob o nº 121.790.318-66, ambos residentes e domiciliados à Rua Antonio Pirola, nº 123, Residencial Alto das Andorinhas, CEP 15046-231, na cidade de São José do Rio Preto - SP; e")

    p_v2 = add_p()
    r = p_v2.add_run("2. WEBER REGINALDO DOS SANTOS")
    r.bold = True
    p_v2.add_run(", brasileiro, empresário/comerciante, portador da Cédula de Identidade RG nº 22.806.400-4 SSP/SP e inscrito no CPF/MF sob o nº 151.990.938-19, casado sob o regime da Comunhão Universal de Bens (conforme Escritura Pública de Pacto Antenupcial lavrada no 4º Tabelião de Notas de São José do Rio Preto/SP, Livro nº 0552, fls. 379) com ")
    r = p_v2.add_run("PATRICIA FERNANDA DE MORAES BALDUINO DOS SANTOS")
    r.bold = True
    p_v2.add_run(", brasileira, portadora da Cédula de Identidade RG nº 30.213.714-2 SSP/SP e inscrita no CPF/MF sob o nº 289.020.208-94, ambos residentes e domiciliados à Rua Direitos Humanos, nº 611, Residencial Ana Célia, CEP 15045-512, na cidade de São José do Rio Preto - SP.")

    # COMPRADORES
    add_clause_header("B – PROMISSÁRIOS COMPRADORES:")
    
    p_c = add_p()
    r = p_c.add_run("JOÃO ROBERTO FRESCHI")
    r.bold = True
    p_c.add_run(", brasileiro, cirurgião-dentista aposentado, portador da Cédula de Identidade RG nº 6.012.357-6 SSP/SP e inscrito no CPF/MF sob o nº 787.804.468-68, e sua esposa ")
    r = p_c.add_run("MARIA GENERCI RIGONATO FRESCHI")
    r.bold = True
    p_c.add_run(", brasileira, professora aposentada / do lar, portadora da Cédula de Identidade RG nº 12.515.252 SSP/SP e inscrita no CPF/MF sob o nº 045.709.208-45, casados entre si sob o regime da Comunhão Universal de Bens na vigência da Lei nº 6.515/77 (conforme Escritura Pública de Pacto Antenupcial lavrada no Cartório de Paraíso - SP, Livro nº 58, fls. 167, devidamente registrada sob o nº 3.266 no 1º Oficial de Registro de Imóveis da Comarca de Catanduva - SP), ambos residentes e domiciliados na Avenida Romeu Strazzi, nº 385, 7º pavimento, Apartamento nº 74, Edifício Ipê da Mata, Bairro Jardim Walkíria / Reserva da Mata, CEP 15085-520, na cidade de São José do Rio Preto - SP.")

    add_p("DECLARAM as partes a sua livre e espontânea vontade de compromissar a venda e compra do imóvel objeto deste contrato, aceitando mutuamente a dação em pagamento parcial por via de permuta imobiliária, obrigando-se por si, seus herdeiros ou sucessores a qualquer título, a cumprirem fielmente as cláusulas e condições a seguir estipuladas:")

    # CLÁUSULA PRIMEIRA
    add_clause_header("CLÁUSULA PRIMEIRA – DO OBJETO DA VENDA")
    add_p("Os PROMITENTES VENDEDORES declaram, sob as penas da Lei, que são legítimos possuidores e proprietários, a justo título, livre e desembaraçado de quaisquer tributos, dívidas, hipotecas ou ônus reais, do imóvel assim caracterizado:")
    p_imovel = add_p()
    p_imovel.paragraph_format.left_indent = Inches(0.2)
    r = p_imovel.add_run("RESIDÊNCIA UNIFAMILIAR DE ALTO PADRÃO")
    r.bold = True
    p_imovel.add_run(", situada no ")
    r = p_imovel.add_run("Condomínio Residencial Parque Buona Vita")
    r.bold = True
    p_imovel.add_run(" (Avenida Percy Gandini, nº 5005, Bairro Parque Residencial Buona Vita), na cidade de São José do Rio Preto - SP, edificada sobre o ")
    r = p_imovel.add_run("Lote nº [____], Quadra nº [____]")
    r.bold = True
    p_imovel.add_run(", com área total de terreno de [____] m² e área construída de [____] m², devidamente cadastrada perante a Prefeitura Municipal de São José do Rio Preto sob a Inscrição Municipal nº [_________________] e registrada sob a Matrícula nº [____________] perante o Oficial de Registro de Imóveis competente da Comarca de São José do Rio Preto - SP.")

    # CLÁUSULA SEGUNDA
    add_clause_header("CLÁUSULA SEGUNDA – DO PREÇO E DA FORMA DE PAGAMENTO")
    add_p("O preço total, certo e ajustado para a presente COMPRA E VENDA é de R$ 2.390.000,00 (dois milhões, trezentos e noventa mil reais), que será pago pelos PROMISSÁRIOS COMPRADORES aos PROMITENTES VENDEDORES rigorosamente na seguinte conformidade:")

    # Parcelas
    p1 = add_p()
    p1.paragraph_format.left_indent = Inches(0.2)
    r = p1.add_run("1. 1ª PARCELA – SINAL E PRINCÍPIO DE PAGAMENTO: ")
    r.bold = True
    p1.add_run("O valor de ")
    r = p1.add_run("R$ 200.000,00 (duzentos mil reais)")
    r.bold = True
    p1.add_run(", pagos no ato da assinatura deste instrumento particular, a título de sinal e princípio de pagamento, através de transferência bancária via PIX/TED em conta corrente bancária indicada de titularidade dos PROMITENTES VENDEDORES, servindo o respectivo comprovante de transação como recibo de quitação desta parcela;")

    p2 = add_p()
    p2.paragraph_format.left_indent = Inches(0.2)
    r = p2.add_run("2. 2ª PARCELA – RECURSOS PRÓPRIOS (60 DIAS): ")
    r.bold = True
    p2.add_run("O valor de ")
    r = p2.add_run("R$ 200.000,00 (duzentos mil reais)")
    r.bold = True
    p2.add_run(", a serem pagos no prazo impreterível de até 60 (sessenta) dias contados da assinatura deste contrato, prazo este ajustado em comum acordo entre as partes para viabilizar o resgate e liquidação de aplicação financeira dos PROMISSÁRIOS COMPRADORES;")

    p3 = add_p()
    p3.paragraph_format.left_indent = Inches(0.2)
    r = p3.add_run("3. 3ª PARCELA – FINANCIAMENTO BANCÁRIO: ")
    r.bold = True
    p3.add_run("O valor de ")
    r = p3.add_run("R$ 380.000,00 (trezentos e oitenta mil reais)")
    r.bold = True
    p3.add_run(", a serem pagos através de recursos obtidos via Financiamento Imobiliário a ser contratado pelos PROMISSÁRIOS COMPRADORES perante Instituição Financeira de sua livre escolha (SFH/SFI), comprometendo-se os PROMITENTES VENDEDORES a fornecer tempestivamente todas as certidões, documentos pessoais e do imóvel que forem exigidos pelo agente financeiro para a competente análise e liberação do crédito;")

    p4 = add_p()
    p4.paragraph_format.left_indent = Inches(0.2)
    r = p4.add_run("4. 4ª PARCELA – DAÇÃO EM PAGAMENTO / PERMUTA IMOBILIÁRIA: ")
    r.bold = True
    p4.add_run("O valor de ")
    r = p4.add_run("R$ 1.610.000,00 (um milhão, seiscentos e dez mil reais)")
    r.bold = True
    p4.add_run(", representados pela dação em pagamento, transferência de domínio, posse e propriedade do seguinte imóvel pertencente aos PROMISSÁRIOS COMPRADORES em favor dos PROMITENTES VENDEDORES:")

    # Descrição completa do Ipê da Mata
    p_permuta = add_p()
    p_permuta.paragraph_format.left_indent = Inches(0.4)
    r = p_permuta.add_run("DO IMÓVEL DADO EM DAÇÃO / PERMUTA: ")
    r.bold = True
    p_permuta.add_run("O ")
    r = p_permuta.add_run("APARTAMENTO RESIDENCIAL Nº 74 (setenta e quatro)")
    r.bold = True
    p_permuta.add_run(", localizado no 7º (sétimo) pavimento do ")
    r = p_permuta.add_run("Condomínio 'IPÊ DA MATA'")
    r.bold = True
    p_permuta.add_run(", situado na Avenida Romeu Strazzi, nº 385, Bairro Jardim Walkíria / Reserva da Mata, na cidade de São José do Rio Preto - SP, contendo a área privativa de ")
    r = p_permuta.add_run("127,27 m²")
    r.bold = True
    p_permuta.add_run("; acrescida de 21,62 m² de área privativa de ")
    r = p_permuta.add_run("02 (duas) vagas de garagem (nºs 88-A e 88-B)")
    r.bold = True
    p_permuta.add_run(", localizadas no pavimento térreo (com 10,81 m² cada uma); área comum de 55,2684 m², perfazendo a área total construída de ")
    r = p_permuta.add_run("204,1584 m²")
    r.bold = True
    p_permuta.add_run(", com fração ideal no terreno e coisas comuns de 1,1829%, devidamente cadastrado na Prefeitura Municipal de São José do Rio Preto sob o ")
    r = p_permuta.add_run("Cadastro nº 419444028")
    r.bold = True
    p_permuta.add_run(" e registrado sob a ")
    r = p_permuta.add_run("Matrícula nº 117.310 perante o 2º Oficial de Registro de Imóveis de São José do Rio Preto - SP")
    r.bold = True
    p_permuta.add_run(", adquirido por força de Escritura Pública lavrada no 2º Tabelião de Notas local (Livro 1210, Páginas 203/204).")

    add_p("Parágrafo Primeiro: O presente contrato é celebrado em caráter irretratável e irrevogável, restaurando-se as partes ao status quo ante exclusivamente na hipótese de rescisão motivada nos termos da Cláusula Nona deste instrumento.")

    add_p("Parágrafo Segundo (Posse e Entrega Simultânea das Chaves): As partes contratantes convencionam que a entrega das chaves e a respectiva imissão na posse de ambos os imóveis (tanto da residência no Condomínio Buona Vita quanto do apartamento nº 74 no Edifício Ipê da Mata) ocorrerão de forma simultânea no prazo de até 60 (sessenta) dias contados a partir da data de assinatura deste contrato, mediante a constatação do adimplemento das obrigações pecuniárias iniciais.")

    p2_3 = add_p()
    p2_3.paragraph_format.left_indent = Inches(0.2)
    r = p2_3.add_run("- Parágrafo Terceiro (Do Prazo de 20 Dias Úteis para Venda do Apartamento e Consolidação da Permuta): ")
    r.bold = True
    p2_3.add_run("Fica expressamente estabelecido e convencionado entre os contratantes que os ")
    r = p2_3.add_run("PROMISSÁRIOS COMPRADORES")
    r.bold = True
    p2_3.add_run(" terão o prazo improrrogável de até ")
    r = p2_3.add_run("20 (vinte) dias úteis")
    r.bold = True
    p2_3.add_run(", contados rigorosamente a partir da data de assinatura deste contrato, para buscar e concretizar a venda do ")
    r = p2_3.add_run("Apartamento nº 74 do Condomínio Edifício Ipê da Mata")
    r.bold = True
    p2_3.add_run(" perante terceiros adquirentes no mercado imobiliário. Caso a referida venda a terceiros não venha a ser efetivamente concluída e formalizada dentro do mencionado prazo estipulado de 20 (vinte) dias úteis, a dação em pagamento / permuta imobiliária se consolidará de pleno direito e em caráter definitivo, irrevogável e irretratável, passando o referido apartamento a ser de legítima, plena e exclusiva propriedade dos ")
    r = p2_3.add_run("PROMITENTES VENDEDORES")
    r.bold = True
    p2_3.add_run(" pelo valor de avaliação de ")
    r = p2_3.add_run("R$ 1.610.000,00 (um milhão, seiscentos e dez mil reais)")
    r.bold = True
    p2_3.add_run(", mantendo-se inalterados o valor total da transação de ")
    r = p2_3.add_run("R$ 2.390.000,00 (dois milhões, trezentos e noventa mil reais)")
    r.bold = True
    p2_3.add_run(" e a obrigação dos vendedores referente à instalação dos aparelhos de ar-condicionado na residência do Buona Vita prevista na Cláusula Quinta deste instrumento.")

    p2_4 = add_p()
    p2_4.paragraph_format.left_indent = Inches(0.2)
    r = p2_4.add_run("- Parágrafo Quarto (Da Venda do Apartamento, Repactuação do Valor para R$ 2.250.000,00 À Vista e Exoneração dos Ares-Condicionados): ")
    r.bold = True
    p2_4.add_run("Na hipótese de os ")
    r = p2_4.add_run("PROMISSÁRIOS COMPRADORES")
    r.bold = True
    p2_4.add_run(" concretizarem com êxito a venda do Apartamento nº 74 do Edifício Ipê da Mata a terceiros dentro do aludido prazo de 20 (vinte) dias úteis, as partes estipulam que vigorarão de forma imediata e automática as seguintes condições especiais:\n")
    p2_4.add_run("a) O valor total da presente negociação da residência unifamiliar no Condomínio Parque Buona Vita será repactuado e passará a ser de ")
    r = p2_4.add_run("R$ 2.250.000,00 (dois milhões, duzentos e cinquenta mil reais)")
    r.bold = True
    p2_4.add_run(";\n")
    p2_4.add_run("b) A forma de pagamento passará a ser na modalidade integralmente ")
    r = p2_4.add_run("À VISTA")
    r.bold = True
    p2_4.add_run(", computando-se e deduzindo-se o sinal de R$ 200.000,00 (duzentos mil reais) pago no ato da assinatura deste instrumento e quitando-se integralmente o saldo remanescente de ")
    r = p2_4.add_run("R$ 2.050.000,00 (dois milhões e cinquenta mil reais)")
    r.bold = True
    p2_4.add_run(" na data da formalização da escritura/contrato de venda do apartamento a terceiros e correspondente liberação financeira dos recursos;\n")
    r = p2_4.add_run("c) Exoneração Total dos Vendedores quanto aos Ares-Condicionados: ")
    r.bold = True
    p2_4.add_run("Ocorrendo a venda do apartamento a terceiros e a fixação do preço com desconto à vista de R$ 2.250.000,00, ")
    r = p2_4.add_run("os PROMITENTES VENDEDORES ficarão formal e expressamente exonerados e desobrigados de adquirir, fornecer e instalar quaisquer aparelhos de ar-condicionado")
    r.bold = True
    p2_4.add_run(" na residência do Condomínio Parque Buona Vita (ficando desobrigados tanto dos 05 aparelhos individuais quanto do ar-condicionado modelo Cassette do espaço gourmet), sendo a residência entregue sem os referidos aparelhos, permanecendo outrossim inalterada a supressão da lareira conforme acordo comercial prévio.")

    # CLÁUSULA TERCEIRA
    add_clause_header("CLÁUSULA TERCEIRA – DAS CERTIDÕES E REGULARIDADE DOCUMENTAL")
    add_p("Os PROMITENTES VENDEDORES e os PROMISSÁRIOS COMPRADORES (em relação ao imóvel dado em dação em pagamento) comprometem-se reciprocamente a apresentar, no prazo improrrogável de até 15 (quinze) dias a contar da assinatura deste contrato, todas as certidões de praxe indispensáveis à comprovação da segurança jurídica do negócio:")
    
    p_cert1 = add_p()
    p_cert1.paragraph_format.left_indent = Inches(0.2)
    r = p_cert1.add_run("a) Certidões Pessoais dos Contratantes: ")
    r.bold = True
    p_cert1.add_run("Certidões dos Distribuidores de Protestos de Títulos (últimos 5 anos); Certidões de Distribuições Cíveis, Família, Criminais e Execuções Fiscais Estaduais (últimos 10 anos); Certidões Cíveis, Criminais e de Execução da Justiça Federal; Certidões Negativas de Débitos Trabalhistas (CNDT) da Justiça do Trabalho; e Certidões Conjuntas Negativas de Débitos Relativos aos Tributos Federais e à Dívida Ativa da União expedidas pela Secretaria da Receita Federal do Brasil / PGFN.")

    p_cert2 = add_p()
    p_cert2.paragraph_format.left_indent = Inches(0.2)
    r = p_cert2.add_run("b) Documentos e Certidões dos Imóveis: ")
    r.bold = True
    p_cert2.add_run("Certidões de Inteiro Teor das Matrículas com Negativa de Ônus Reais e Alienações atualizadas; Certidões Negativas de Débitos Tributários Municipais (IPTU); Certidões de Débitos Condominiais e de Taxas de Associação de Moradores com expressa declaração de quitação emitida pela respectiva administradora; e Certidões de Inexistência de Débitos junto às concessionárias locais prestadoras de serviços públicos (SEMAE e CPFL Paulista).")

    add_p("Parágrafo Único: Caso seja apontada certidão positiva cujas informações esclarecedoras (certidões de objeto e pé) revelem gravames, dívidas ou contingências patrimoniais insuperáveis capazes de colocar em risco a segurança do negócio ou a transferência do domínio, a parte adquirente poderá rejeitá-las e pleitear a rescisão sem a incidência de qualquer penalidade ou multa contratual.")

    # CLÁUSULA QUARTA
    add_clause_header("CLÁUSULA QUARTA – DA VISTORIA E DO ESTADO DOS IMÓVEIS")
    add_p("As partes declaram expressamente que realizaram prévia vistoria física em ambos os imóveis (tanto na casa do Condomínio Buona Vita quanto no apartamento nº 74 do Edifício Ipê da Mata), conhecendo perfeitamente suas dimensões, divisões, estado de conservação, instalações elétricas, hidráulicas e acabamentos, aceitando-os nas condições exatas em que se encontram nesta data, ressalvadas as obrigações e benfeitorias especificadas na Cláusula Quinta.")

    # CLÁUSULA QUINTA
    add_clause_header("CLÁUSULA QUINTA – DO ACORDO COMERCIAL DE BENFEITORIAS E EQUIPAMENTOS")
    add_p("Para composição harmônica dos valores, especificações técnicas e equilíbrio comercial entre os imóveis transacionados, os contratantes ajustam as seguintes condições comerciais quanto às benfeitorias e equipamentos:")

    p_eq1 = add_p()
    p_eq1.paragraph_format.left_indent = Inches(0.2)
    r = p_eq1.add_run("- Parágrafo Primeiro (Equipamentos no Apartamento Ipê da Mata - Permuta): ")
    r.bold = True
    p_eq1.add_run("Na hipótese de consolidação da permuta tratada na Cláusula Segunda deste contrato, fica expressamente convencionado que permanecerão instalados e serão entregues pelos PROMISSÁRIOS COMPRADORES no Apartamento nº 74 do Condomínio Edifício Ipê da Mata o total de ")
    r = p_eq1.add_run("05 (cinco) aparelhos de ar-condicionado")
    r.bold = True
    p_eq1.add_run(" no estado de conservação e perfeito funcionamento em que se encontram, com todas as suas tubulações, fiações, condensadoras, evaporadoras e respectivos controles remotos, os quais são transferidos aos PROMITENTES VENDEDORES sem qualquer custo adicional.")

    p_eq2 = add_p()
    p_eq2.paragraph_format.left_indent = Inches(0.2)
    r = p_eq2.add_run("- Parágrafo Segundo (Instalação dos Aparelhos de Ar-Condicionado na Casa do Buona Vita pelos Vendedores): ")
    r.bold = True
    p_eq2.add_run("Ressalvada a hipótese de venda à vista tratada no Parágrafo Quarto da Cláusula Segunda, os PROMITENTES VENDEDORES assumem a responsabilidade e obrigação de adquirir, fornecer e instalar, por sua conta e ônus exclusivo, na residência do Condomínio Parque Buona Vita, o total de ")
    r = p_eq2.add_run("05 (cinco) aparelhos de ar-condicionado")
    r.bold = True
    p_eq2.add_run(", novos e em perfeito estado de funcionamento, rigorosamente nas seguintes potências e especificações técnicas:\n")
    r = p_eq2.add_run("I. 03 (três) aparelhos de ar-condicionado com capacidade de 12.000 BTUs;\n")
    r.bold = True
    r = p_eq2.add_run("II. 01 (um) aparelho de ar-condicionado com capacidade de 18.000 BTUs;\n")
    r.bold = True
    r = p_eq2.add_run("III. 01 (um) aparelho de ar-condicionado com capacidade de 24.000 BTUs;\n")
    r.bold = True
    r = p_eq2.add_run("IV. ALÉM de 01 (um) aparelho de ar-condicionado modelo Cassette com capacidade de 60.000 BTUs, ")
    r.bold = True
    p_eq2.add_run("a ser devidamente instalado, embutido no teto e testado no ambiente do Espaço Gourmet.")

    p_eq3 = add_p()
    p_eq3.paragraph_format.left_indent = Inches(0.2)
    r = p_eq3.add_run("- Parágrafo Terceiro (Supressão da Lareira e Substituição pelo Ar-Condicionado Cassette de 60.000 BTUs): ")
    r.bold = True
    p_eq3.add_run("Fica expressamente salientado e acordado que na residência do Condomínio Parque Buona Vita ")
    r = p_eq3.add_run("NÃO será instalada a lareira")
    r.bold = True
    p_eq3.add_run(" que constava originalmente do projeto da área de lazer (fire place), por ter sido objeto de expresso e livre acordo comercial entre as partes, tendo sido a referida lareira substituída em definitivo pelo fornecimento e instalação do aparelho de ")
    r = p_eq3.add_run("ar-condicionado modelo Cassette de 60.000 BTUs no Espaço Gourmet")
    r.bold = True
    p_eq3.add_run(" a cargo exclusivo dos PROMITENTES VENDEDORES, dando os contratantes mútua, plena e irrevogável quitação quanto a este item.")

    p_eq4 = add_p()
    p_eq4.paragraph_format.left_indent = Inches(0.2)
    r = p_eq4.add_run("- Parágrafo Quarto (Da Desoneração dos Vendedores em Caso de Venda do Apartamento à Vista): ")
    r.bold = True
    p_eq4.add_run("Conforme estabelecido no Parágrafo Quarto da Cláusula Segunda, caso os PROMISSÁRIOS COMPRADORES concretizem a venda do apartamento a terceiros dentro do prazo de 20 (vinte) dias úteis e o negócio da residência no Condomínio Buona Vita seja liquidado pelo valor de R$ 2.250.000,00 à vista, as obrigações dos PROMITENTES VENDEDORES estipuladas no Parágrafo Segundo desta Cláusula (instalação dos 05 aparelhos de ar-condicionado e do Cassette de 60.000 BTUs) ")
    r = p_eq4.add_run("ficarão plena e automaticamente extintas e sem efeito")
    r.bold = True
    p_eq4.add_run(", não sendo instalado nenhum aparelho de ar-condicionado pelos vendedores na referida residência, mantendo-se igualmente a supressão da lareira.")

    # CLÁUSULA SEXTA
    add_clause_header("CLÁUSULA SEXTA – DA IRRETRATABILIDADE E ADJUDICAÇÃO COMPULSÓRIA")
    add_p("O presente contrato obriga os contratantes de modo irrevogável e irretratável. Em caso de recusa injustificada ou omissão de quaisquer das partes em outorgar as competentes Escrituras Públicas Definitivas de Venda e Compra e de Permuta após o cumprimento das obrigações pactuadas e liberação do financiamento bancário, o presente instrumento conferirá o direito à Adjudicação Compulsória, servindo como título executivo e hábil nos termos do art. 22 do Decreto-Lei nº 58/1937, da Lei nº 649/1949, do art. 1.418 do Código Civil e das normas processuais aplicáveis.")

    # CLÁUSULA SÉTIMA
    add_clause_header("CLÁUSULA SÉTIMA – DA ESCRITURAÇÃO E DO REGISTRO")
    add_p("As Escrituras Públicas Definitivas de Compra e Venda e de Dação em Pagamento / Permuta serão outorgadas e lavradas imediatamente após a liberação do crédito referente ao Financiamento Bancário pactuado no item 3 da Cláusula Segunda, correndo as respectivas despesas nos termos estabelecidos na Cláusula Décima Primeira deste contrato.")

    p_p1_c7 = add_p()
    p_p1_c7.paragraph_format.left_indent = Inches(0.2)
    r = p_p1_c7.add_run("- Parágrafo Primeiro (Da Procuração Pública do Imóvel em Permuta): ")
    r.bold = True
    p_p1_c7.add_run("Especificamente em relação ao imóvel dado em permuta (Apartamento nº 74 do Condomínio Ipê da Mata - Matrícula nº 117.310 do 2º ORI), ")
    r = p_p1_c7.add_run("decorrido o prazo de 20 (vinte) dias úteis previsto no Parágrafo Terceiro da Cláusula Segunda sem a concretização da venda a terceiros, consolidando-se a permuta")
    r.bold = True
    p_p1_c7.add_run(", os PROMISSÁRIOS COMPRADORES obrigam-se formalmente a outorgar em favor dos PROMITENTES VENDEDORES (ou de pessoa física/jurídica por estes expressamente indicada), no prazo improrrogável de até 05 (cinco) dias úteis subsequentes, Instrumento Público de Procuração lavrado em Tabelionato de Notas competente.")

    # CLÁUSULA OITAVA
    add_clause_header("CLÁUSULA OITAVA – DA CLÁUSULA PENAL COMPENSATÓRIA")
    add_p("A parte que inadimplir, descumprir ou obstaculizar a execução de qualquer das cláusulas deste contrato ficará sujeita ao pagamento de multa penal de 10% (dez por cento) incidente sobre o valor total do negócio (R$ 2.390.000,00), devidamente atualizado pelo índice oficial (IPCA/IBGE) desde a assinatura até a efetiva liquidação, sem prejuízo da responsabilidade pelas perdas e danos e fixação de honorários advocatícios em 20% (vinte por cento) em caso de procedimento judicial.")

    # CLÁUSULA NONA
    add_clause_header("CLÁUSULA NONA – DA RESCISÃO CONTRATUAL")
    add_p("O contrato somente poderá ser rescindido em caso de infração contratual insanável ou inadimplemento das obrigações financeiras por qualquer dos polos, ensejando a aplicação da Cláusula Penal e o retorno imediato ao estado anterior, com as devidas compensações legais.")

    # CLÁUSULA DÉCIMA
    add_clause_header("CLÁUSULA DÉCIMA – DA INTERMEDIAÇÃO IMOBILIÁRIA")
    p_corretagem = add_p()
    p_corretagem.add_run("A comissão de corretagem pela intermediação imobiliária, fixada no percentual legal e usual de ")
    r = p_corretagem.add_run("5% (cinco por cento)")
    r.bold = True
    p_corretagem.add_run(" sobre o valor da transação, é devida e será paga ao corretor de imóveis credenciado ")
    r = p_corretagem.add_run("VALTEIR DE OLIVEIRA – CRECI/SP 214072F")
    r.bold = True
    p_corretagem.add_run(" (titular de ")
    r = p_corretagem.add_run("Valteir Imóveis - CNPJ nº 50.185.716/0001-89")
    r.bold = True
    p_corretagem.add_run(", dados bancários: Banco C6 S.A., Chave PIX CNPJ 50.185.716/0001-89), na proporção e condições ajustadas entre as partes e o profissional intermediador.")

    # CLÁUSULA DÉCIMA PRIMEIRA
    add_clause_header("CLÁUSULA DÉCIMA PRIMEIRA – DOS TRIBUTOS, TAXAS CONDOMINIAIS E DESPESAS CARTORÁRIAS")
    add_p("Todos os tributos municipais (IPTU), contribuições condominiais ordinárias e extraordinárias, taxas associativas e faturas de consumo de água (SEMAE), energia elétrica (CPFL) e gás incidentes sobre cada um dos imóveis serão de inteira responsabilidade dos respectivos transmitentes até a data da efetiva imissão na posse (entrega das chaves), momento a partir do qual a responsabilidade passa a ser integral dos adquirentes.")
    
    add_p("Parágrafo Único: Todas as despesas cartorárias necessárias à lavratura da Escritura Pública, recolhimento do Imposto de Transmissão de Bens Imóveis (ITBI) e emolumentos de Registro de Imóveis relativos à aquisição do imóvel residencial no Condomínio Buona Vita serão de responsabilidade exclusiva dos PROMISSÁRIOS COMPRADORES. Da mesma forma, as despesas notariais, tributárias (ITBI) e registrais relativas à transferência do Apartamento nº 74 do Condomínio Ipê da Mata serão de inteira responsabilidade dos PROMITENTES VENDEDORES (novos titulares do apartamento). As partes obrigam-se a transferir a titularidade de contas e cadastros de IPTU e concessionárias no prazo de até 30 (trinta) dias após a posse.")

    # CLÁUSULA DÉCIMA SEGUNDA - REFORÇADA COM PROTEÇÃO ESPECÍFICA AOS VENDEDORES EM CASO DE MORTE DOS COMPRADORES
    add_clause_header("CLÁUSULA DÉCIMA SEGUNDA – DA GARANTIA ESPECÍFICA DA PERMUTA, PROTEÇÃO EM CASO DE FALECIMENTO DOS COMPRADORES E VINCULAÇÃO DE HERDEIROS")
    add_p("Considerando que a dação em pagamento/permuta do Apartamento nº 74 do Edifício Ipê da Mata constitui fração substancial (R$ 1.610.000,00) e indissociável do equilíbrio econômico-financeiro da compra e venda pactuada, e considerando a existência de filhos e herdeiros necessários dos PROMISSÁRIOS COMPRADORES, as partes ajustam com força cogente, vinculante e irretratável as seguintes garantias em favor dos PROMITENTES VENDEDORES:")

    p_gar1 = add_p()
    p_gar1.paragraph_format.left_indent = Inches(0.2)
    r = p_gar1.add_run("1. Ato Jurídico Perfeito e Consolidação de Direitos: ")
    r.bold = True
    p_gar1.add_run("A alienação por via de permuta/dação em pagamento do Apartamento nº 74 do Edifício Ipê da Mata constitui negócio jurídico oneroso, comutativo, perfeito e acabado na data da assinatura deste instrumento, pelo qual os PROMISSÁRIOS COMPRADORES transferem em caráter definitivo aos PROMITENTES VENDEDORES todos os direitos aquisitivos, de posse e de propriedade;")

    p_gar2 = add_p()
    p_gar2.paragraph_format.left_indent = Inches(0.2)
    r = p_gar2.add_run("2. Incomunicabilidade com a Partilha do Espólio em Caso de Morte: ")
    r.bold = True
    p_gar2.add_run("Na eventualidade de falecimento, interdição civil ou incapacidade superveniente de um ou de ambos os PROMISSÁRIOS COMPRADORES antes da lavratura ou registro da competente Escritura Pública Definitiva de Permuta, o imóvel (Apartamento nº 74 do Edifício Ipê da Mata) ")
    r = p_gar2.add_run("NÃO integrará nem sofrerá bloqueio na partilha de bens entre os filhos e herdeiros necessários")
    r.bold = True
    p_gar2.add_run(", por se tratar de bem previamente alienado a título oneroso e com transmissão de posse operada de boa-fé, constituindo obrigação contratual infungível a ser cumprida integralmente pelo espólio;")

    p_gar3 = add_p()
    p_gar3.paragraph_format.left_indent = Inches(0.2)
    r = p_gar3.add_run("3. Obrigação Expressa de Requerimento de Alvará Judicial pelos Herdeiros: ")
    r.bold = True
    p_gar3.add_run("Ocorrendo o evento de morte, o(a) cônjuge sobrevivente, inventariante, testamenteiro, bem como todos os filhos e herdeiros necessários dos PROMISSÁRIOS COMPRADORES, assumem a obrigação formal, legal e expressa de ")
    r = p_gar3.add_run("pleitear imediatamente nos autos do competente processo de Inventário (judicial ou por Escritura Pública Extrajudicial) a expedição de Alvará Judicial específico ou a outorga da Escritura Definitiva de Permuta/Venda e Compra")
    r.bold = True
    p_gar3.add_run(" em favor dos PROMITENTES VENDEDORES (Ademir Venturini e Weber Reginaldo dos Santos e respectivas esposas) ou a quem estes expressamente indicarem, no prazo máximo de 60 (sessenta) dias a contar da abertura da sucessão;")

    p_gar4 = add_p()
    p_gar4.paragraph_format.left_indent = Inches(0.2)
    r = p_gar4.add_run("4. Renúncia a Oposições e Retenção: ")
    r.bold = True
    p_gar4.add_run("Fica expressamente vedado ao espólio, cessionários ou herdeiros dos PROMISSÁRIOS COMPRADORES suscitar qualquer embargo, retenção, reivindicação ou oposição material ou processual relativamente à posse e domínio do Apartamento nº 74 do Ipê da Mata, servindo o presente contrato como título hábil para a imissão de posse compulsória e para a Ação de Adjudicação Compulsória em face do Espólio, nos termos dos arts. 1.418 do Código Civil, 22 do Decreto-Lei nº 58/1937 e normas correlatas.")

    # CLÁUSULA DÉCIMA TERCEIRA
    add_clause_header("CLÁUSULA DÉCIMA TERCEIRA – DA EVICÇÃO E INEXISTÊNCIA DE ÔNUS")
    add_p("Os PROMITENTES VENDEDORES e os PROMISSÁRIOS COMPRADORES declaram formalmente, sob as penas da Lei civil e penal, que os respectivos imóveis transacionados acham-se inteiramente livres e desembaraçados de ônus reais, hipotecas, penhoras, ações reipersecutórias, execuções judiciais ou gravames de qualquer espécie, responsabilizando-se mutuamente pela evicção de direito nos exatos termos dos arts. 447 a 457 do Código Civil Brasileiro.")

    # CLÁUSULA DÉCIMA QUARTA
    add_clause_header("CLÁUSULA DÉCIMA QUARTA – DO FORO")
    add_p("Para dirimir quaisquer controvérsias, divergências ou conflitos oriundos da execução e interpretação do presente instrumento, as partes elegem expressamente o Foro da Comarca de São José do Rio Preto - Estado de São Paulo, renunciando expressamente a qualquer outro foro por mais privilegiado ou especial que seja.")

    add_p("E, por estarem assim justos e contratados, assinam o presente instrumento em 03 (três) vias de igual teor e forma, na presença de 02 (duas) testemunhas instrumentárias abaixo qualificadas, para que produza todos os seus efeitos jurídicos e legais.")

    # Data
    add_p("São José do Rio Preto - SP, [Dia] de [Mês] de 2026.", align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=16, space_after=24)

    # Assinaturas dos Vendedores
    add_clause_header("PROMITENTES VENDEDORES:")
    
    add_p("____________________________________________________\nADEMIR VENTURINI\nCPF: 133.470.468-62", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=18)
    
    add_p("____________________________________________________\nMARILAINE OLIVEIRA SILVA VENTURINI\nCPF: 121.790.318-66", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=18)

    add_p("____________________________________________________\nWEBER REGINALDO DOS SANTOS\nCPF: 151.990.938-19", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=18)

    add_p("____________________________________________________\nPATRICIA FERNANDA DE MORAES BALDUINO DOS SANTOS\nCPF: 289.020.208-94", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=24)

    # Assinaturas dos Compradores
    add_clause_header("PROMISSÁRIOS COMPRADORES:")
    
    add_p("____________________________________________________\nJOÃO ROBERTO FRESCHI\nCPF: 787.804.468-68", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=18)

    add_p("____________________________________________________\nMARIA GENERCI RIGONATO FRESCHI\nCPF: 045.709.208-45", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=24)

    # Testemunhas
    add_clause_header("TESTEMUNHAS:")

    add_p("1. ________________________________________________\nNome:\nCPF:", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=14)
    add_p("2. ________________________________________________\nNome:\nCPF:", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=14)

    # Termo Opcional de Ciência e Anuência dos Filhos / Herdeiros Necessários
    add_clause_header("TERMO ANEXO DE CIÊNCIA E ANUÊNCIA DOS HERDEIROS NECESSÁRIOS (OPCIONAL/FACULTATIVO):")
    add_p("Na qualidade de filhos e herdeiros necessários dos PROMISSÁRIOS COMPRADORES JOÃO ROBERTO FRESCHI e MARIA GENERCI RIGONATO FRESCHI, declaramos ter pleno e irrevogável conhecimento dos termos e condições do presente negócio imobiliário e da dação em pagamento/permuta do Apartamento nº 74 do Condomínio Ipê da Mata em favor dos PROMITENTES VENDEDORES, ratificando integralmente todas as disposições da Cláusula Décima Segunda e comprometendo-nos a não opor qualquer resistência à lavratura da escritura definitiva ou ao cumprimento de alvará:")

    add_p("____________________________________________________\nNome do Herdeiro(a):\nCPF:", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=16)
    add_p("____________________________________________________\nNome do Herdeiro(a):\nCPF:", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=16)
    add_p("____________________________________________________\nNome do Herdeiro(a):\nCPF:", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=16)

    # Salvar nos destinos
    # Destino 1: Workspace do projeto (OneDrive)
    workspace_file = r"c:\Users\valte\OneDrive\Área de Trabalho\Projeto - Valteir Imoveis\Contrato_Compromisso_Compra_e_Venda_Buona_Vita_Ipe_da_Mata.docx"
    doc.save(workspace_file)
    print(f"Salvo no Workspace: {workspace_file}")

    # Destino 2: Pasta oficial do Buona Vita / Reginaldo
    target_dir = r"C:\Users\valte\OneDrive\Área de Trabalho\Negocios Imobiliários\Condomínio Buona Vita\Reginaldo"
    if not os.path.exists(target_dir):
        target_dir = r"D:\Negocios Imobiliários\Condomínio Buona Vita\Reginaldo"

    if os.path.exists(target_dir):
        target_updated = os.path.join(target_dir, "Contrato_Compromisso_Compra_e_Venda_Buona_Vita_Ipe_da_Mata.docx")
        doc.save(target_updated)
        print(f"Salvo na pasta de destino: {target_updated}")

if __name__ == "__main__":
    create_contract()
