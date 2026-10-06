import os
import copy
import shutil
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

FOLDER = r"C:\Users\valte\OneDrive\Área de Trabalho\Negocios Imobiliários\Condomínio Buona Vita\Reginaldo"
PROJ = r"c:\Users\valte\OneDrive\Área de Trabalho\Projeto - Valteir Imoveis"

src = [f for f in os.listdir(FOLDER) if f.endswith("(Versão Final - 15 dias).docx")][0]
src_path = os.path.join(FOLDER, src)
shutil.copy2(src_path, os.path.join(FOLDER, "Contrato Buona Vita_BACKUP6_antes_ar_condicionado.docx"))

doc = docx.Document(src_path)


def find(prefix):
    return next(p for p in doc.paragraphs if p.text.strip().startswith(prefix))


def rewrite(p, parts):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for text, bold in parts:
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.bold = bold


# --- Parágrafo Primeiro: equipamentos que ficam no apartamento Ipê da Mata ---
p1 = find("- Parágrafo Primeiro (Equipamentos no Apartamento")
rewrite(p1, [
    ("- Parágrafo Primeiro (Equipamentos que Permanecerão no Apartamento do Edifício Ipê da Mata - Permuta): ", True),
    ("Na hipótese de consolidação da permuta tratada na Cláusula Segunda deste contrato, fica expressamente convencionado que permanecerão instalados e serão entregues pelos PROMISSÁRIOS COMPRADORES no ", False),
    ("Apartamento nº 74 do Edifício Ipê da Mata", True),
    (", sem qualquer custo adicional aos PROMITENTES VENDEDORES, os seguintes ", False),
    ("05 (cinco) aparelhos de ar-condicionado", True),
    (", no estado de conservação e perfeito funcionamento em que se encontram, com todas as suas tubulações, fiações, condensadoras, evaporadoras e respectivos controles remotos:\n", False),
    ("I. 03 (três) aparelhos de ar-condicionado com capacidade de 12.000 BTUs;\n", True),
    ("II. 01 (um) aparelho de ar-condicionado com capacidade de 18.000 BTUs;\n", True),
    ("III. 01 (um) aparelho de ar-condicionado com capacidade de 24.000 BTUs.", True),
])

# --- Parágrafo Segundo: instalação na casa do Buona Vita, após a 2ª parcela ---
p2 = find("- Parágrafo Segundo (Instalação dos Aparelhos")
rewrite(p2, [
    ("- Parágrafo Segundo (Instalação dos Aparelhos de Ar-Condicionado na Casa do Buona Vita pelos Vendedores): ", True),
    ("Em contrapartida aos equipamentos que permanecerão no apartamento, descritos no Parágrafo Primeiro, e ressalvada a hipótese de venda à vista tratada no Parágrafo Quarto da Cláusula Segunda, os PROMITENTES VENDEDORES assumem a responsabilidade e obrigação de adquirir, fornecer e instalar, por sua conta e ônus exclusivo, na residência do Condomínio Parque Buona Vita:\n", False),
    ("I. 05 (cinco) aparelhos de ar-condicionado novos, nas mesmas potências dos equipamentos que permanecerão no apartamento (03 de 12.000 BTUs, 01 de 18.000 BTUs e 01 de 24.000 BTUs);\n", True),
    ("II. ALÉM de 01 (um) aparelho de ar-condicionado modelo Cassette com capacidade de 60.000 BTUs, ", True),
    ("a ser devidamente instalado, embutido no teto e testado no ambiente do Espaço Gourmet.\n", False),
    ("Prazo de Instalação: ", True),
    ("A instalação de todos os aparelhos de ar-condicionado previstos neste Parágrafo será realizada pelos PROMITENTES VENDEDORES ", False),
    ("somente após o efetivo recebimento da 2ª Parcela (R$ 200.000,00 – item 2 da Cláusula Segunda)", True),
    (", devendo estar concluída até a data da entrega das chaves e imissão na posse.", False),
])


# --- Dividir parágrafos com quebras de linha (evita espaçamento esticado no justificado) ---
def split_at_breaks(p):
    body_p = p._p
    if not body_p.findall(".//" + qn("w:br")):
        return 0
    pPr = body_p.find(qn("w:pPr"))
    segments = [[]]
    for r in list(body_p.findall(qn("w:r"))):
        rPr = r.find(qn("w:rPr"))

        def new_run():
            nr = OxmlElement("w:r")
            if rPr is not None:
                nr.append(copy.deepcopy(rPr))
            return nr

        cur = new_run()
        for child in r:
            if child.tag == qn("w:rPr"):
                continue
            if child.tag == qn("w:br") and child.get(qn("w:type")) in (None, "textWrapping"):
                segments[-1].append(cur)
                segments.append([])
                cur = new_run()
            else:
                cur.append(copy.deepcopy(child))
        segments[-1].append(cur)
        body_p.remove(r)

    for nr in segments[0]:
        body_p.append(nr)
    prev = body_p
    for seg in segments[1:]:
        np_ = OxmlElement("w:p")
        if pPr is not None:
            np_.append(copy.deepcopy(pPr))
        for nr in seg:
            np_.append(nr)
        prev.addnext(np_)
        prev = np_
    return len(segments) - 1


total = 0
for p in list(doc.paragraphs):
    if p.alignment is not None and "JUSTIFY" in str(p.alignment):
        total += split_at_breaks(p)
print(f"Quebras de linha convertidas em parágrafos: {total}")

# Salvar
try:
    doc.save(src_path)
    print(f"Salvo em: {src}")
except PermissionError:
    alt = src.replace(".docx", " - ar condicionado.docx")
    doc.save(os.path.join(FOLDER, alt))
    print(f"Arquivo aberto no Word. Salvo como: {alt}")

for t in [os.path.join(FOLDER, "Contrato_Compromisso_Compra_e_Venda_Buona_Vita_Ipe_da_Mata.docx"),
          os.path.join(PROJ, "Contrato_Compromisso_Compra_e_Venda_Buona_Vita_Ipe_da_Mata.docx")]:
    try:
        doc.save(t)
    except PermissionError:
        print(f"Espelho aberto no Word: {t}")

# Conferência
d = docx.Document(src_path) if os.path.exists(src_path) else doc
start = next(i for i, p in enumerate(d.paragraphs) if p.text.strip().upper().startswith("CLÁUSULA QUINTA"))
for p in d.paragraphs[start:start + 14]:
    print("  >", p.text)
