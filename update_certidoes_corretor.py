import os
import shutil
import docx

FOLDER = r"C:\Users\valte\OneDrive\Área de Trabalho\Negocios Imobiliários\Condomínio Buona Vita\Reginaldo"
PROJ = r"c:\Users\valte\OneDrive\Área de Trabalho\Projeto - Valteir Imoveis"

src = [f for f in os.listdir(FOLDER) if f.endswith("(Versão Final - 15 dias).docx")][0]
src_path = os.path.join(FOLDER, src)
shutil.copy2(src_path, os.path.join(FOLDER, "Contrato Buona Vita_BACKUP5_antes_certidoes.docx"))

doc = docx.Document(src_path)

idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().upper().startswith("CLÁUSULA TERCEIRA"))
target = doc.paragraphs[idx + 1]
assert "certidões de praxe" in target.text, target.text

NEW_TEXT = (
    "As certidões de praxe indispensáveis à comprovação da idoneidade patrimonial e plena higidez do negócio, "
    "relativas aos PROMITENTES VENDEDORES, aos PROMISSÁRIOS COMPRADORES e a ambos os imóveis transacionados "
    "(inclusive o imóvel dado em permuta), serão providenciadas pelo corretor de imóveis intermediador "
    "VALTEIR DE OLIVEIRA – CRECI/SP 214072F, no prazo improrrogável de até 15 (quinze) dias contados da "
    "assinatura deste contrato, obrigando-se os PROMITENTES VENDEDORES e os PROMISSÁRIOS COMPRADORES a "
    "fornecer-lhe, de imediato, todos os documentos, dados pessoais e autorizações necessários à sua obtenção, "
    "compreendendo as seguintes certidões:"
)

runs = target.runs
runs[0].text = NEW_TEXT
runs[0].bold = None
for r in runs[1:]:
    r._r.getparent().remove(r._r)
print("Novo caput:\n", target.text)

try:
    doc.save(src_path)
    print(f"Salvo em: {src}")
except PermissionError:
    alt = src.replace(".docx", " - certidoes.docx")
    doc.save(os.path.join(FOLDER, alt))
    print(f"Arquivo aberto no Word. Salvo como: {alt}")

for t in [os.path.join(FOLDER, "Contrato_Compromisso_Compra_e_Venda_Buona_Vita_Ipe_da_Mata.docx"),
          os.path.join(PROJ, "Contrato_Compromisso_Compra_e_Venda_Buona_Vita_Ipe_da_Mata.docx")]:
    try:
        doc.save(t)
    except PermissionError:
        print(f"Espelho aberto no Word: {t}")
print("Espelhos atualizados.")
