import os
import re
import shutil
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

FOLDER = r"C:\Users\valte\OneDrive\Área de Trabalho\Negocios Imobiliários\Condomínio Buona Vita\Reginaldo"
PROJ = r"c:\Users\valte\OneDrive\Área de Trabalho\Projeto - Valteir Imoveis"

src = [f for f in os.listdir(FOLDER) if "(Atualizado Completo)" in f and f.endswith(".docx")][0]
src_path = os.path.join(FOLDER, src)
original = [f for f in os.listdir(FOLDER) if f.startswith("Contrato Compromisso") and f.endswith(".docx")
            and "BACKUP" not in f and "Atualizado" not in f][0]
original_path = os.path.join(FOLDER, original)

shutil.copy2(src_path, os.path.join(FOLDER, "Contrato Buona Vita_BACKUP3_antes_15dias.docx"))

doc = docx.Document(src_path)

# 1. Substituir 20 -> 15 dias úteis (run a run)
REPL = [
    ("20 (vinte) dias úteis", "15 (quinze) dias úteis"),
    ("20 (vinte) Dias Úteis", "15 (quinze) Dias Úteis"),
    ("20 Dias Úteis", "15 Dias Úteis"),
    ("20 dias úteis", "15 dias úteis"),
]
count = 0
for p in doc.paragraphs:
    for r in p.runs:
        t = r.text
        for a, b in REPL:
            if a in t:
                count += t.count(a)
                t = t.replace(a, b)
        r.text = t
print(f"Substituições 20 -> 15 dias úteis: {count}")

# Verificação de sobras (texto partido entre runs)
for i, p in enumerate(doc.paragraphs):
    if re.search(r"20 \(vinte\) dias|20 Dias Úteis|20 dias úteis", p.text, re.I):
        print(f"ATENÇÃO - sobra em P{i}: {p.text[:120]}")

# 2. Evitar o espaçamento esticado antes de quebras de linha em parágrafos justificados
settings = doc.settings.element
compat = settings.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat")
    settings.append(compat)
if compat.find(qn("w:doNotExpandShiftReturn")) is None:
    compat.insert(0, OxmlElement("w:doNotExpandShiftReturn"))

# 3. Comissão de R$ 100.000,00 na venda à vista
idx = None
for i, p in enumerate(doc.paragraphs):
    if "CLÁUSULA DÉCIMA" in p.text.upper() and "INTERMEDIA" in p.text.upper():
        idx = i
        break
anchor = doc.paragraphs[idx + 1]  # caput da cláusula de comissão

new = OxmlElement("w:p")
anchor._p.addnext(new)
para = Paragraph(new, anchor._parent)
para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
para.paragraph_format.left_indent = Inches(0.2)
para.paragraph_format.space_after = Pt(6)

def run(text, bold=False):
    r = para.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.bold = bold

run("- Parágrafo Único (Comissão na Hipótese de Venda à Vista): ", True)
run("Ocorrendo a hipótese prevista no Parágrafo Quarto da Cláusula Segunda, ou seja, concretizada a venda do Apartamento nº 74 do Edifício Ipê da Mata a terceiros dentro do prazo de 15 (quinze) dias úteis e repactuado o negócio para pagamento à vista no valor de R$ 2.250.000,00, a comissão de corretagem devida ao corretor ")
run("VALTEIR DE OLIVEIRA – CRECI/SP 214072F", True)
run(" passará a ser fixa no valor de ")
run("R$ 100.000,00 (cem mil reais)", True)
run(", em substituição ao critério percentual previsto no caput desta Cláusula, mantida a mesma responsabilidade pelo pagamento, devendo ser quitada na data da liquidação do saldo à vista, mediante transferência bancária (PIX/TED) para a Chave PIX CNPJ 50.185.716/0001-89 (Valteir Imóveis - Banco C6 S.A.).")

# 4. Salvar
out_name = "Contrato Compromisso Compra e Venda - Buona Vita - Reginaldo x João Roberto (Versão Final - 15 dias).docx"
doc.save(os.path.join(FOLDER, out_name))
print(f"Salvo: {out_name}")
for target in [os.path.join(FOLDER, "Contrato_Compromisso_Compra_e_Venda_Buona_Vita_Ipe_da_Mata.docx"),
               os.path.join(PROJ, "Contrato_Compromisso_Compra_e_Venda_Buona_Vita_Ipe_da_Mata.docx")]:
    try:
        doc.save(target)
        print(f"Espelho salvo: {target}")
    except PermissionError:
        print(f"Espelho aberto no Word - não salvo: {target}")

# Verificação final
d = docx.Document(os.path.join(FOLDER, out_name))
for i, p in enumerate(d.paragraphs):
    if "15 (quinze)" in p.text or "100.000,00" in p.text:
        print(f"[{i}] {p.text[:160]}")
